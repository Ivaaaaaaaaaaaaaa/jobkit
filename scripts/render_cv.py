#!/usr/bin/env python3
"""Render a CV to DOCX (and PDF) from data, not by editing Word files.

    me/cv/base.yaml                  full CV: all roles, bullets, skills
    data/applications/<id>/cv.yaml   overlay: only what changes per vacancy
    jk cv [overlay.yaml] [--pdf] | --selftest

Overlay fields (all optional): title, summary, experience {company: [bullet indexes]},
role_titles {company: title from title_options}, skills_order [groups], hide [companies], file_name.
"""
from __future__ import annotations

import copy
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from paths import ME, DATA  # noqa: E402
BASE = ME / "cv" / "base.yaml"


def load_yaml(p: Path) -> dict:
    import yaml
    return yaml.safe_load(p.read_text(encoding="utf-8")) or {}


def merge(base: dict, ov: dict) -> dict:
    cv = copy.deepcopy(base)
    if ov.get("title"):
        cv["title"] = ov["title"]
    if ov.get("summary"):
        cv["summary"] = ov["summary"]
    hide = set(ov.get("hide") or [])
    picks = ov.get("experience") or {}
    titles = ov.get("role_titles") or {}
    roles = []
    for r in cv.get("experience", []):
        if r.get("company") in hide:
            continue
        if r.get("company") in picks:
            idx = picks[r["company"]]
            r["bullets"] = [r["bullets"][i] for i in idx if i < len(r.get("bullets", []))]
        if r.get("company") in titles:
            allowed = [r.get("role")] + (r.get("title_options") or [])
            if titles[r["company"]] not in allowed:
                raise SystemExit(f"title {titles[r['company']]!r} for {r['company']} is not in title_options "
                                 f"in base.yaml — invented titles are not allowed")
            r["role"] = titles[r["company"]]
        roles.append(r)
    cv["experience"] = roles
    order = ov.get("skills_order")
    if order and cv.get("skills"):
        rank = {k: i for i, k in enumerate(order)}
        cv["skills"] = sorted(cv["skills"], key=lambda s: rank.get(s.get("group"), 99))
    return cv


def render(cv: dict, out: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm, RGBColor

    doc = Document()
    sec = doc.sections[0]
    if str(cv.get("page", "A4")).upper() == "A4":      # US: page: Letter
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    sec.left_margin = sec.right_margin = Cm(1.6)
    width = sec.page_width - sec.left_margin - sec.right_margin
    st = doc.styles["Normal"]
    st.font.name = cv.get("font", "Calibri")
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(0)

    def para(text="", size=10, bold=False, italic=False, after=0, color=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text)
            r.font.size, r.bold, r.italic = Pt(size), bold, italic
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        return p

    def link(p, text, url, size=9):
        rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                               is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        for tag, val in (("w:color", "1F4E99"), ("w:u", "single"), ("w:sz", str(size * 2))):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            rpr.append(e)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        p._p.append(h)

    def heading(text):
        p = para(text.upper(), size=10.5, bold=True, after=2)
        p.paragraph_format.space_before = Pt(7)
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "999999")):
            b.set(qn(k), v)
        pbdr.append(b)
        p._p.get_or_add_pPr().append(pbdr)

    def two_sides(left, right, bold=False, italic=False):
        p = para(after=0)
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left)
        r.bold, r.italic, r.font.size = bold, italic, Pt(10)
        r = p.add_run("\t" + (right or ""))
        r.italic, r.font.size = italic, Pt(10)
        return p

    h = cv.get("header", {})
    para(h.get("name", ""), size=16, bold=True)
    contacts = [x for x in (cv.get("title"), h.get("location"), h.get("phone"), h.get("email")) if x]
    p = para("  ·  ".join(contacts), size=9, after=0)
    for name, url in (h.get("links") or {}).items():
        p.add_run("  ·  ").font.size = Pt(9)
        link(p, name, url)

    if cv.get("summary"):
        heading("Summary")
        para(" ".join(str(cv["summary"]).split()), size=10)

    if cv.get("experience"):
        heading("Experience")
        for r in cv["experience"]:
            left = r.get("company", "") + (f"  —  {r['descriptor']}" if r.get("descriptor") else "")
            two_sides(left, r.get("location", ""), bold=True).paragraph_format.space_before = Pt(4)
            two_sides(r.get("role", ""), r.get("dates", ""), italic=True)
            for b in r.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(0)
                run = bp.add_run(" ".join(str(b).split()))
                run.font.size = Pt(10)

    if cv.get("projects"):
        heading("Projects")
        for pr in cv["projects"]:
            p = para(after=0)
            r = p.add_run(pr.get("name", "") + ": ")
            r.bold, r.font.size = True, Pt(10)
            p.add_run(" ".join(str(pr.get("text", "")).split())).font.size = Pt(10)
            if pr.get("url"):
                p.add_run("  ")
                link(p, "link", pr["url"])

    if cv.get("skills"):
        heading("Skills")
        for s in cv["skills"]:
            p = para(after=0)
            r = p.add_run(f"{s.get('group', '')}: ")
            r.bold, r.font.size = True, Pt(10)
            p.add_run(", ".join(s.get("items", []))).font.size = Pt(10)

    if cv.get("education"):
        heading("Education")
        for e in cv["education"]:
            two_sides(e.get("school", "") + (f" — {e['degree']}" if e.get("degree") else ""), e.get("dates", ""))

    extra = [x for x in (cv.get("languages"), cv.get("work_authorization")) if x]
    if extra:
        para("  ·  ".join(extra), size=9.5, after=0).paragraph_format.space_before = Pt(5)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def to_pdf(docx: Path) -> Path | None:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if Path(soffice).exists() or shutil.which("soffice"):
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=False, capture_output=True)
    else:
        try:
            from docx2pdf import convert  # uses MS Word on Mac/Windows
            convert(str(docx), str(docx.with_suffix(".pdf")))
        except Exception:
            print("no PDF: install LibreOffice (brew install --cask libreoffice) or docx2pdf")
            return None
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def page_count_hint(cv: dict) -> int:
    """Rough line count: > ~58 lines almost surely spills onto page two."""
    n = 6 + (len(str(cv.get("summary", ""))) // 110 + 2)
    for r in cv.get("experience", []):
        n += 3 + sum(len(str(b)) // 105 + 1 for b in r.get("bullets", []))
    n += 2 + len(cv.get("skills", [])) + 2 + len(cv.get("education", [])) + 2 * len(cv.get("projects", []))
    return n


def selftest() -> int:
    base = {"header": {"name": "Test Person", "email": "t@example.com", "links": {"LinkedIn": "https://x.io"}},
            "title": "Product Manager", "summary": "Base summary.",
            "experience": [{"company": "A", "role": "PM", "title_options": ["Product Lead"], "dates": "2024–now",
                            "bullets": ["one", "two", "three"]},
                           {"company": "B", "role": "APM", "dates": "2022–2024", "bullets": ["x"]}],
            "skills": [{"group": "Data", "items": ["SQL"]}, {"group": "Product", "items": ["Discovery"]}],
            "education": [{"school": "Uni", "degree": "BSc", "dates": "2018–2022"}]}
    cv = merge(base, {"title": "Growth PM", "experience": {"A": [2, 0]}, "role_titles": {"A": "Product Lead"},
                      "skills_order": ["Product"], "hide": ["B"]})
    assert cv["experience"][0]["bullets"] == ["three", "one"] and len(cv["experience"]) == 1
    assert cv["skills"][0]["group"] == "Product" and cv["experience"][0]["role"] == "Product Lead"
    try:
        merge(base, {"role_titles": {"A": "CEO"}})
        raise AssertionError("invented title accepted")
    except SystemExit:
        pass
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        assert render(cv, Path(d) / "t.docx").stat().st_size > 5000
    print("render_cv selftest: ok")
    return 0


def main() -> int:
    if "--selftest" in sys.argv:
        return selftest()
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not BASE.exists():
        sys.exit("no me/cv/base.yaml — run /setup or /cv-interview")
    base = load_yaml(BASE)
    if args:
        ov_path = Path(args[0])
        ov = load_yaml(ov_path)
        cv = merge(base, ov)
        name = ov.get("file_name") or f"{base.get('header', {}).get('name', 'CV')} CV"
        out = ov_path.parent / f"{name}.docx"
    else:
        cv = merge(base, {})
        out = DATA / f"{base.get('header', {}).get('name', 'CV')} CV.docx"
    render(cv, out)
    lines = page_count_hint(cv)
    from events import log
    log("cv_render", kind="tailored" if args and "applications" in args[0] else ("variant" if args else "base"), lines=lines)
    print(f"DOCX: {out}")
    if lines > 58:
        print(f"⚠ ~{lines} lines — likely over one page; drop bullets via overlay -> experience")
    if "--pdf" in sys.argv:
        pdf = to_pdf(out)
        if pdf:
            print(f"PDF:  {pdf}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
