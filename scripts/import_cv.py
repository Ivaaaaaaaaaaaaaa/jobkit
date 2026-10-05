#!/usr/bin/env python3
"""Turn a CV or LinkedIn profile export into plain text the agent can structure.

    jk import <file.pdf|.docx|.txt|.md> [--as cv|linkedin]
LinkedIn: Profile -> More -> Save to PDF, then import that PDF.
Writes data/import/<as>-<name>.txt and prints a short summary (not the whole text — the agent
reads the file once). Text stays local.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from paths import DATA  # noqa: E402


def extract(p: Path) -> str:
    ext = p.suffix.lower()
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    return p.read_text(encoding="utf-8", errors="replace")


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print(__doc__)
        return 2
    kind = "linkedin" if "linkedin" in " ".join(sys.argv).lower() else "cv"
    src = Path(args[0]).expanduser()
    if not src.exists():
        print(f"no file: {src}")
        return 1
    text = re.sub(r"\n{3,}", "\n\n", extract(src)).strip()
    out = DATA / "import" / f"{kind}-{re.sub(r'[^a-z0-9]+', '-', src.stem.lower()).strip('-')}.txt"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    years = sorted(set(re.findall(r"\b(?:19|20)\d\d\b", text)))
    nums = len(re.findall(r"\d+\s?%|\b\d+[kKmM]\b|[$€£]\s?\d", text))
    print(f"{out} · {len(text.split())} words · years {years[0] if years else '?'}–{years[-1] if years else '?'} · metrics found {nums}")
    if len(text.split()) < 80:
        print("warning: very little text — scanned PDF? ask for DOCX or paste the text")
    return 0


if __name__ == "__main__":
    sys.exit(main())
